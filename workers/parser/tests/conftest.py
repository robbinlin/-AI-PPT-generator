"""Programmatic fixtures: a 3-slide PPTX and a 3-page PDF, plus JSON Schema validators."""

from __future__ import annotations

import json
import shutil
from pathlib import Path

import pymupdf
import pytest
from jsonschema import Draft7Validator
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
PARSER_DIR = HERE.parent
REPO_ROOT = PARSER_DIR.parent.parent
SCHEMA_DIR = REPO_ROOT / "packages" / "shared" / "schemas"


def _validator(name: str) -> Draft7Validator:
    schema = json.loads((SCHEMA_DIR / f"{name}.json").read_text(encoding="utf-8"))
    return Draft7Validator(schema)


@pytest.fixture(scope="session")
def manifest_validator() -> Draft7Validator:
    return _validator("parser-manifest")


@pytest.fixture(scope="session")
def content_validator() -> Draft7Validator:
    return _validator("slide-content")


def validate_output(out_dir: Path, manifest_validator: Draft7Validator, content_validator: Draft7Validator) -> dict:
    manifest = json.loads((out_dir / "manifest.json").read_text(encoding="utf-8"))
    manifest_validator.validate(manifest)
    assert manifest["pageCount"] == len(manifest["pages"])
    for entry in manifest["pages"]:
        for key in ("imagePath", "thumbnailPath", "contentPath"):
            path = out_dir / entry[key]
            assert path.is_file(), f"missing {key}: {path}"
        content = json.loads((out_dir / entry["contentPath"]).read_text(encoding="utf-8"))
        content_validator.validate(content)
        assert content["pageNumber"] == entry["pageNumber"]
        with Image.open(out_dir / entry["imagePath"]) as img:
            assert (img.width, img.height) == (entry["imageWidthPx"], entry["imageHeightPx"])
    return manifest


@pytest.fixture(scope="session")
def sample_png(tmp_path_factory: pytest.TempPathFactory) -> Path:
    path = tmp_path_factory.mktemp("img") / "pic.png"
    Image.new("RGB", (200, 120), "#3b82f6").save(path)
    return path


@pytest.fixture(scope="session")
def sample_pptx(tmp_path_factory: pytest.TempPathFactory, sample_png: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "運動生理學 第三章"
    prs.core_properties.author = "測試作者"

    # Slide 1: title + subtitle
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "最大攝氧量 VO₂max 的測量"
    s1.placeholders[1].text = "運動生理學 第三章"
    s1.notes_slide.notes_text_frame.text = "開場先問學生有沒有做過跑步機測試。"

    # Slide 2: title + two-level bullets + picture
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "測量方法"
    body = s2.placeholders[1].text_frame
    body.text = "直接測量法"
    p = body.add_paragraph()
    p.text = "氣體分析儀"
    p.level = 1
    p = body.add_paragraph()
    p.text = "間接推估法"
    p.level = 0
    s2.shapes.add_picture(str(sample_png), Inches(8), Inches(2), width=Inches(4))

    # Slide 3: title + table + chart
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "常模比較"
    rows, cols = 3, 3
    table = s3.shapes.add_table(rows, cols, Inches(0.5), Inches(1.5), Inches(6), Inches(2)).table
    cells = [["組別", "平均", "SD"], ["介入組", "48.2", "4.1"], ["對照組", "44.9", "3.8"]]
    for r in range(rows):
        for c in range(cols):
            table.cell(r, c).text = cells[r][c]
            table.cell(r, c).text_frame.paragraphs[0].font.size = Pt(14)
    chart_data = CategoryChartData()
    chart_data.categories = ["前測", "後測"]
    chart_data.add_series("介入組", (44.0, 48.2))
    chart_data.add_series("對照組", (44.5, 44.9))
    chart = s3.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(7), Inches(1.5), Inches(5.5), Inches(4), chart_data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = "VO2max 變化"

    path = tmp_path_factory.mktemp("pptx") / "deck.pptx"
    prs.save(path)
    return path


@pytest.fixture(scope="session")
def sample_pdf(tmp_path_factory: pytest.TempPathFactory) -> Path:
    doc = pymupdf.open()
    page = doc.new_page(width=960, height=540)
    page.insert_text((60, 90), "Chapter 3: Measuring VO2max", fontsize=32, fontname="helv")
    page.insert_text((60, 160), "Direct measurement uses gas analysis.", fontsize=16, fontname="helv")
    page.insert_text((60, 190), "Indirect estimation uses heart rate.", fontsize=16, fontname="helv")
    page = doc.new_page(width=960, height=540)
    page.insert_text((60, 90), "Normative values", fontsize=32, fontname="helv")
    page.insert_text((60, 160), "Intervention group 48.2 vs control 44.9", fontsize=16, fontname="helv")
    # Page 3: scanned-style page with only a drawing, no text layer
    page = doc.new_page(width=960, height=540)
    page.draw_rect(pymupdf.Rect(100, 100, 800, 400), color=(0, 0, 0), fill=(0.8, 0.8, 0.8))
    doc.set_metadata({"title": "Exercise Physiology", "author": "Tester"})
    path = tmp_path_factory.mktemp("pdf") / "deck.pdf"
    doc.save(path)
    doc.close()
    return path


@pytest.fixture(scope="session")
def sample_ppt(tmp_path_factory: pytest.TempPathFactory, sample_pptx: Path) -> Path | None:
    """Legacy .ppt produced by LibreOffice from the PPTX fixture (None when soffice is absent)."""
    from lvs_parser.convert import ConvertError, soffice_binary, soffice_convert

    if soffice_binary() is None:
        return None
    out = tmp_path_factory.mktemp("ppt")
    try:
        return soffice_convert(sample_pptx, out, "ppt", timeout=180)
    except ConvertError:
        return None


@pytest.fixture
def no_soffice(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("SOFFICE_BIN", "/nonexistent/soffice-missing")
    assert shutil.which("/nonexistent/soffice-missing") is None
