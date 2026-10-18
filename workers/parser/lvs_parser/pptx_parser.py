"""PPTX structured-content extraction with python-pptx."""

from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.shapes.base import BaseShape
from pptx.util import Emu

from .manifest import SlideContent, bbox, text_block

EMU_PER_PT = 12700.0
TITLE_PLACEHOLDERS = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
SUBTITLE_PLACEHOLDERS = {PP_PLACEHOLDER.SUBTITLE}


def _pt(value: int | None) -> float:
    return float(value or 0) / EMU_PER_PT


def _shape_bbox(shape: BaseShape) -> dict[str, float] | None:
    if shape.left is None or shape.top is None:
        return None
    return bbox(_pt(shape.left), _pt(shape.top), _pt(shape.width), _pt(shape.height))


def _placeholder_type(shape: BaseShape):
    try:
        if shape.is_placeholder:
            return shape.placeholder_format.type
    except (AttributeError, KeyError, ValueError):
        return None
    return None


def _iter_shapes(shapes: Iterable[BaseShape]) -> Iterable[BaseShape]:
    """Depth-first over shapes, flattening groups."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _text_frame_blocks(shape: BaseShape, role: str, box: dict[str, float] | None) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    for paragraph in shape.text_frame.paragraphs:
        text = "".join(run.text for run in paragraph.runs).strip() or paragraph.text.strip()
        if not text:
            continue
        blocks.append(text_block(text, role, box, paragraph.level))
    return blocks


def _chart_summary(shape: BaseShape) -> dict[str, Any]:
    chart = shape.chart
    info: dict[str, Any] = {}
    try:
        info["chartType"] = str(chart.chart_type).split(".")[-1].split(" ")[0]
    except (AttributeError, ValueError):
        pass
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            info["title"] = chart.chart_title.text_frame.text.strip()
    except (AttributeError, ValueError):
        pass
    parts: list[str] = []
    try:
        for plot in chart.plots:
            categories = [str(c) for c in plot.categories]
            if categories:
                parts.append("categories: " + ", ".join(categories[:20]))
            for series in plot.series:
                values = ", ".join(f"{v:g}" if isinstance(v, (int, float)) else str(v) for v in list(series.values)[:20])
                parts.append(f"{series.name}: {values}")
    except (AttributeError, ValueError, TypeError):
        pass
    if parts:
        info["dataSummary"] = "; ".join(parts)
    box = _shape_bbox(shape)
    if box:
        info["bbox"] = box
    return info


def _table_rows(shape: BaseShape) -> list[list[str]]:
    return [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]


def _notes_text(slide) -> str | None:
    if not slide.has_notes_slide:
        return None
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return None
    text = frame.text.strip()
    return text or None


def parse_slide(slide, page_number: int, width_pt: float, height_pt: float, fidelity: str) -> SlideContent:
    content = SlideContent(page_number=page_number, width_pt=width_pt, height_pt=height_pt, render_fidelity=fidelity)
    for shape in _iter_shapes(slide.shapes):
        box = _shape_bbox(shape)
        ph_type = _placeholder_type(shape)
        if shape.has_text_frame and shape.text_frame.text.strip():
            if ph_type in TITLE_PLACEHOLDERS and content.title is None:
                content.title = shape.text_frame.text.strip().replace("\n", " ")
                content.text_blocks.append(text_block(content.title, "title", box))
                continue
            role = "subtitle" if ph_type in SUBTITLE_PLACEHOLDERS else ("bullet" if ph_type is not None else "body")
            content.text_blocks.extend(_text_frame_blocks(shape, role, box))
        if getattr(shape, "has_table", False) and shape.has_table:
            table: dict[str, Any] = {"rows": _table_rows(shape)}
            if box:
                table["bbox"] = box
            content.tables.append(table)
        if getattr(shape, "has_chart", False) and shape.has_chart:
            content.charts.append(_chart_summary(shape))
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image: dict[str, Any] = {}
            if box:
                image["bbox"] = box
            alt = (shape._element.xpath("./p:nvPicPr/p:cNvPr/@descr") or [None])[0]
            if alt:
                image["altText"] = str(alt)
            content.images.append(image)
    content.notes = _notes_text(slide)
    return content


def parse_pptx(pptx_path: Path, fidelity: str = "converted") -> tuple[list[SlideContent], dict[str, str]]:
    prs = Presentation(str(pptx_path))
    width_pt = _pt(prs.slide_width or Emu(12192000))
    height_pt = _pt(prs.slide_height or Emu(6858000))
    pages = [parse_slide(slide, index, width_pt, height_pt, fidelity) for index, slide in enumerate(prs.slides, start=1)]
    meta: dict[str, str] = {}
    props = prs.core_properties
    if props.title:
        meta["title"] = props.title
    if props.author:
        meta["author"] = props.author
    return pages, meta
