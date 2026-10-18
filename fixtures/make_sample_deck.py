"""Builds fixtures/sample-exercise-physiology.pptx: a 6-slide zh-TW teaching deck with
statistics terms, a table, speaker notes and a chart, used for end-to-end validation."""
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("sample-exercise-physiology.pptx")


def add_title_body(prs, title, bullets, notes=None):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title
    body = slide.placeholders[1].text_frame
    body.clear()
    for i, (level, text) in enumerate(bullets):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = text
        p.level = level
        p.font.size = Pt(24 if level == 0 else 20)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = "高強度間歇訓練對最大攝氧量的影響"
    s.placeholders[1].text = "運動生理學 第八週 | 授課：林教授"
    s.notes_slide.notes_text_frame.text = "開場先回顧上週的能量系統，再帶入 HIIT 的定義。"

    add_title_body(prs, "學習目標", [
        (0, "說明 VO₂max 的定義與測量方式"),
        (0, "比較 HIIT 與 MICT 的生理適應差異"),
        (0, "解讀介入研究中的 ANOVA 結果與 95% 信賴區間"),
    ], notes="強調第三點是本週的評量重點。")

    add_title_body(prs, "VO₂max 的定義", [
        (0, "最大攝氧量：單位時間內身體能攝取並利用的最大氧氣量"),
        (1, "絕對值：L/min；相對值：mL/kg/min"),
        (0, "Fick 方程式：VO₂ = Q × (a-v)O₂ diff"),
        (1, "Q 為心輸出量，(a-v)O₂ diff 為動靜脈血氧差"),
    ], notes="可提問：為什麼耐力選手的相對值通常高於絕對值的差異？")

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "介入研究：HIIT vs MICT（8 週）"
    rows = [
        ["組別", "n", "前測 VO₂max", "後測 VO₂max", "變化量", "p 值"],
        ["HIIT 組", "20", "42.1 ± 4.3", "47.8 ± 4.6", "+5.7", "< .001"],
        ["MICT 組", "20", "41.8 ± 4.1", "44.9 ± 4.4", "+3.1", ".004"],
        ["對照組", "18", "42.0 ± 4.5", "42.3 ± 4.7", "+0.3", ".61"],
    ]
    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(0.8), Inches(1.8), Inches(11.5), Inches(2.5)).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            table.cell(r, c).text = val
    slide.notes_slide.notes_text_frame.text = "單位為 mL/kg/min。Gibala et al., 2012 的設計類似，可一併提。"

    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "組間差異（Two-way ANOVA）"
    cd = CategoryChartData()
    cd.categories = ["前測", "第 4 週", "第 8 週"]
    cd.add_series("HIIT", (42.1, 45.0, 47.8))
    cd.add_series("MICT", (41.8, 43.2, 44.9))
    cd.add_series("對照", (42.0, 42.1, 42.3))
    slide.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(0.8), Inches(1.6), Inches(11.5), Inches(5), cd)
    slide.notes_slide.notes_text_frame.text = "交互作用顯著 F(4, 110) = 6.21, p < .001, η² = .18。95% CI 要解釋成區間估計而非機率。"

    add_title_body(prs, "小結與下週預告", [
        (0, "HIIT 在相同時間下帶來較大的 VO₂max 提升"),
        (0, "效果量與信賴區間比單看 p 值更能說明臨床意義"),
        (0, "下週：乳酸閾值與訓練處方"),
    ], notes="提醒作業：閱讀 Gibala 2012 並摘要研究設計。")

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
